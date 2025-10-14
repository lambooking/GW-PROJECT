"""
DOCX文档批注器实现 - 修复版
✅ 修复：避免使用 doc.paragraphs.index() 导致的错误
"""

import logging
import shutil
from pathlib import Path
from typing import List, Union, Optional
from datetime import datetime

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor, Pt
    from docx.enum.text import WD_COLOR_INDEX
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

from .base import BaseAnnotator
from .schemas import Annotation, AnnotationSeverity

logger = logging.getLogger(__name__)


class DocxAnnotator(BaseAnnotator):
    """DOCX文档批注器"""
    
    def __init__(self):
        """初始化DOCX批注器"""
        super().__init__()
        self.supported_extensions = ['.docx']
        
        if not PYTHON_DOCX_AVAILABLE:
            raise ImportError("需要安装 python-docx: pip install python-docx")
        
        self._comment_id = 0
    
    def annotate(
        self,
        input_file: Union[str, Path],
        annotations: List[Annotation],
        output_file: Union[str, Path]
    ) -> Path:
        """在DOCX文档上添加批注"""
        input_file = self._validate_file(input_file)
        output_file = Path(output_file)
        
        logger.info(f"开始为DOCX文档添加批注: {input_file}")
        self._log_annotation_stats(annotations)
        
        # 复制原文件
        shutil.copy2(input_file, output_file)
        
        # 打开文档
        doc = Document(output_file)
        
        # ✅ 改进策略：统一在文档末尾添加批注汇总
        # 原因：python-docx对内联批注支持有限，且在段落间插入容易出错
        logger.info("采用末尾汇总方式添加批注")
        
        # 高亮文档中的相关段落
        highlighted_count = self._highlight_relevant_paragraphs(doc, annotations)
        logger.info(f"高亮了 {highlighted_count} 个相关段落")
        
        # 在文档末尾添加批注汇总
        self._add_annotations_summary(doc, annotations)
        
        # 保存文档
        doc.save(output_file)
        
        logger.info(f"成功添加 {len(annotations)} 条批注，已保存到: {output_file}")
        return output_file
    
    def _highlight_relevant_paragraphs(
        self, 
        doc: Document, 
        annotations: List[Annotation]
    ) -> int:
        """
        高亮文档中与批注相关的段落
        ✅ 改进：直接遍历所有段落，避免索引问题
        """
        highlighted_count = 0
        
        for annotation in annotations:
            if not annotation.text_snippet:
                continue
            
            # 搜索匹配的文本
            snippet = annotation.text_snippet.strip()[:80]  # 取前80字符
            
            # 遍历所有段落
            for para in doc.paragraphs:
                if snippet in para.text:
                    # 高亮该段落
                    self._highlight_paragraph(para, annotation.severity)
                    highlighted_count += 1
                    break  # 找到第一个匹配就停止
        
        return highlighted_count
    
    def _highlight_paragraph(
        self, 
        paragraph, 
        severity: AnnotationSeverity
    ) -> None:
        """高亮段落"""
        try:
            # 根据严重程度选择高亮颜色
            if severity == AnnotationSeverity.CRITICAL:
                highlight_color = WD_COLOR_INDEX.RED
            elif severity == AnnotationSeverity.WARNING:
                highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                highlight_color = WD_COLOR_INDEX.TURQUOISE
            
            # 高亮段落中的所有文本
            for run in paragraph.runs:
                run.font.highlight_color = highlight_color
        except Exception as e:
            logger.debug(f"高亮段落失败: {e}")
    
    def _add_annotations_summary(
        self,
        doc: Document,
        annotations: List[Annotation]
    ) -> None:
        """
        ✅ 改进：在文档末尾添加完整的批注汇总
        这是最稳定可靠的方式
        """
        
        if not annotations:
            return
        
        # 添加分页符
        doc.add_page_break()
        
        # 添加标题
        title = doc.add_paragraph()
        title_run = title.add_run("📋 评分批注汇总")
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(0, 51, 102)
        
        # 添加分隔线
        doc.add_paragraph("=" * 80)
        
        # 添加生成时间
        timestamp = doc.add_paragraph()
        timestamp_run = timestamp.add_run(
            f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        timestamp_run.font.size = Pt(9)
        timestamp_run.font.color.rgb = RGBColor(128, 128, 128)
        
        doc.add_paragraph()
        
        # 按严重程度分组
        critical = [a for a in annotations if a.severity == AnnotationSeverity.CRITICAL]
        warning = [a for a in annotations if a.severity == AnnotationSeverity.WARNING]
        info = [a for a in annotations if a.severity == AnnotationSeverity.INFO]
        
        # 添加统计信息
        stats_para = doc.add_paragraph()
        stats_para.add_run("批注统计\n").bold = True
        
        stats_para.add_run(f"• 总批注数: {len(annotations)} 条\n")
        
        critical_run = stats_para.add_run(f"• 严重问题: {len(critical)} 条\n")
        critical_run.font.color.rgb = RGBColor(220, 20, 60)
        critical_run.bold = True
        
        warning_run = stats_para.add_run(f"• 一般问题: {len(warning)} 条\n")
        warning_run.font.color.rgb = RGBColor(255, 140, 0)
        
        info_run = stats_para.add_run(f"• 建议改进: {len(info)} 条")
        info_run.font.color.rgb = RGBColor(30, 144, 255)
        
        doc.add_paragraph()
        doc.add_paragraph("=" * 80)
        
        # 添加详细批注列表
        severity_groups = [
            ("🔴 严重问题", critical, RGBColor(220, 20, 60)),
            ("🟡 一般问题", warning, RGBColor(255, 140, 0)),
            ("🔵 建议改进", info, RGBColor(30, 144, 255))
        ]
        
        for severity_title, severity_list, color in severity_groups:
            if not severity_list:
                continue
            
            # 添加分组标题
            doc.add_paragraph()
            section_para = doc.add_paragraph()
            section_run = section_para.add_run(f"\n{severity_title} ({len(severity_list)} 条)")
            section_run.bold = True
            section_run.font.size = Pt(14)
            section_run.font.color.rgb = color
            
            doc.add_paragraph("-" * 80)
            
            # 添加每条批注
            for idx, ann in enumerate(severity_list, 1):
                # 批注编号和位置
                header_para = doc.add_paragraph()
                header_run = header_para.add_run(
                    f"\n批注 {idx}: {ann.location}"
                )
                header_run.bold = True
                header_run.font.size = Pt(11)
                
                # 评分项
                item_para = doc.add_paragraph()
                item_para.add_run("评分项: ").bold = True
                item_para.add_run(ann.score_item)
                
                # 扣分信息
                score_para = doc.add_paragraph()
                score_run = score_para.add_run(
                    f"扣分: {ann.score_lost:.1f} / {ann.max_score:.1f} 分"
                )
                score_run.font.color.rgb = color
                score_run.bold = True
                
                # 问题描述
                content_para = doc.add_paragraph()
                content_para.add_run("问题描述:\n").bold = True
                content_para.add_run(ann.content)
                
                # 修改建议
                if ann.suggestion:
                    suggestion_para = doc.add_paragraph()
                    suggestion_para.add_run("修改建议:\n").bold = True
                    suggestion_run = suggestion_para.add_run(ann.suggestion)
                    suggestion_run.font.color.rgb = RGBColor(0, 128, 0)
                
                # 相关文本片段
                if ann.text_snippet:
                    snippet_para = doc.add_paragraph()
                    snippet_para.add_run("相关内容:\n").bold = True
                    snippet_run = snippet_para.add_run(
                        ann.text_snippet[:150] + "..." 
                        if len(ann.text_snippet) > 150 
                        else ann.text_snippet
                    )
                    snippet_run.font.size = Pt(9)
                    snippet_run.font.color.rgb = RGBColor(100, 100, 100)
                    snippet_run.font.italic = True
                
                # 添加分隔线
                doc.add_paragraph("-" * 80)
        
        # 添加结束标记
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer_run = footer.add_run("--- 批注汇总结束 ---")
        footer_run.font.size = Pt(10)
        footer_run.font.color.rgb = RGBColor(128, 128, 128)
        footer_run.italic = True
    
    def _find_target_paragraph(
        self, 
        doc: Document, 
        annotation: Annotation
    ):
        """
        查找目标段落
        ✅ 改进：增强匹配逻辑
        """
        
        # 策略1: 如果有文本片段，精确匹配
        if annotation.text_snippet:
            snippet = annotation.text_snippet.strip()[:100]
            for para in doc.paragraphs:
                if snippet in para.text:
                    return para
        
        # 策略2: 根据章节名称匹配
        if annotation.section_name:
            for para in doc.paragraphs:
                # 章节名称可能在段落开头
                if para.text.strip().startswith(annotation.section_name):
                    return para
                # 或者包含在段落中
                if annotation.section_name in para.text:
                    return para
        
        # 策略3: 根据位置描述模糊匹配
        if annotation.location:
            location_keywords = annotation.location.split()
            for para in doc.paragraphs:
                # 段落必须包含至少一个关键词
                if any(keyword in para.text for keyword in location_keywords if len(keyword) > 2):
                    return para
        
        return None
    
    def _group_annotations(self, annotations: List[Annotation]) -> dict:
        """按位置分组批注"""
        grouped = {}
        for ann in annotations:
            key = (ann.page_number, ann.section_name or ann.location)
            if key not in grouped:
                grouped[key] = []
            grouped[key].append(ann)
        return grouped