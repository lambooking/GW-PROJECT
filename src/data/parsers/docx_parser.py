"""
DOCX document parser implementation.
"""

import base64
import logging
import subprocess
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import List, Optional
from io import BytesIO

try:
    from docx import Document as DocxDocument
    from docx.shared import Inches
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

from .base import BaseDocumentParser
from ..schemas import DocumentInfo, TextContent, Table, Image, StandardizedDocument
from ...core.exceptions import DocumentProcessingError


logger = logging.getLogger(__name__)

# 检查 antiword 是否可用
ANTIWORD_AVAILABLE = shutil.which('antiword') is not None
if ANTIWORD_AVAILABLE:
    logger.info("antiword 命令可用，支持 .doc 文件解析")
else:
    logger.warning("antiword 命令不可用，.doc 文件将无法解析（请安装: apt-get install antiword 或 yum install antiword）")


class DocxDocumentParser(BaseDocumentParser):
    """DOCX document parser using python-docx."""
    
    def __init__(self, enable_ocr: bool = True, enable_images: bool = True):
        """Initialize DOCX parser."""
        super().__init__(enable_ocr, enable_images)
        
        if not PYTHON_DOCX_AVAILABLE:
            raise DocumentProcessingError("python-docx not available. Install with: pip install python-docx")
    
    def supports_format(self, file_path: Path) -> bool:
        """Check if parser supports DOCX format."""
        return file_path.suffix.lower() in ['.docx', '.doc']
    
    def _is_zip_file(self, file_path: Path) -> bool:
        """
        检测文件是否为 ZIP 格式（.docx 实际上是 ZIP 压缩包）
        
        Args:
            file_path: 文件路径
            
        Returns:
            True 如果文件是 ZIP 格式，False 否则
        """
        try:
            with open(file_path, 'rb') as f:
                # ZIP 文件的魔数是 PK (0x504B)
                magic = f.read(2)
                return magic == b'PK'
        except Exception as e:
            logger.debug(f"检测文件格式时出错: {e}")
            return False
    
    def _parse_document(self, file_path: Path) -> StandardizedDocument:
        """Parse DOCX or DOC document."""
        self._validate_file_size(file_path)
        
        # 智能文件格式检测
        is_zip_format = self._is_zip_file(file_path)
        file_suffix = file_path.suffix.lower()
        
        # 情况1: .doc 后缀但实际是 ZIP 格式（错误的扩展名，实际是 .docx）
        if file_suffix == '.doc' and is_zip_format:
            logger.warning(f"文件 {file_path.name} 后缀为 .doc 但实际是 .docx 格式，将使用 python-docx 解析")
            # 继续使用 python-docx 解析
        
        # 情况2: .doc 后缀且不是 ZIP 格式（真正的旧版 .doc）
        elif file_suffix == '.doc' and not is_zip_format:
            return self._parse_doc_with_antiword(file_path)
        
        # 情况3: .docx 后缀或检测为 ZIP 格式
        # 对 .docx 文件使用 python-docx
        try:
            doc = DocxDocument(file_path)
            
            # Extract document info
            document_info = self._extract_document_info(file_path, doc)
            
            # Extract text content
            text_content = self._extract_text_content(doc)
            
            # Extract tables
            tables = self._extract_tables(doc)
            
            # Extract images if enabled  
            images = []
            if self.enable_images:
                images = self._extract_images(doc)
            
            return StandardizedDocument(
                document_info=document_info,
                text_content=text_content,
                tables=tables,
                images=images,
                raw_content=self._get_raw_text(text_content)
            )
            
        except Exception as e:
            raise DocumentProcessingError(f"DOCX parsing error: {e}")
    
    def _extract_document_info(self, file_path: Path, doc: DocxDocument) -> DocumentInfo:
        """Extract document metadata."""
        # Estimate page count (rough approximation)
        total_chars = sum(len(p.text) for p in doc.paragraphs)
        estimated_pages = max(1, total_chars // 2000)  # Rough estimate
        
        return DocumentInfo(
            file_name=file_path.name,
            file_path=str(file_path.absolute()),
            total_pages=estimated_pages,
            file_size=self._get_file_size(file_path),
            processing_timestamp=datetime.now()
        )
    
    def _extract_text_content(self, doc: DocxDocument) -> List[TextContent]:
        """Extract text content from DOCX."""
        text_content = []
        current_page = 1  # We can't easily determine actual page numbers in DOCX
        
        for para_idx, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                # Determine section type and hierarchy
                section_type, hierarchy_level = self._analyze_paragraph(paragraph)
                
                content = TextContent(
                    section_type=section_type,
                    content=paragraph.text.strip(),
                    page_number=current_page,
                    hierarchy_level=hierarchy_level,
                    word_count=len(paragraph.text.strip())
                )
                
                text_content.append(content)
                
                # Simple page estimation (every ~20 paragraphs = 1 page)
                if (para_idx + 1) % 20 == 0:
                    current_page += 1
        
        return text_content
    
    def _extract_tables(self, doc: DocxDocument) -> List[Table]:
        """Extract tables from DOCX."""
        tables = []
        
        for table_idx, table in enumerate(doc.tables):
            try:
                # Extract headers (first row)
                headers = []
                if len(table.rows) > 0:
                    headers = [cell.text.strip() for cell in table.rows[0].cells]
                
                # Extract data (remaining rows)
                data = []
                for row in table.rows[1:]:
                    row_data = [cell.text.strip() for cell in row.cells]
                    data.append(row_data)
                
                table_obj = Table(
                    table_id=f"table_{table_idx + 1}",
                    headers=headers,
                    data=data,
                    page_number=1  # We can't determine actual page in DOCX easily
                )
                
                tables.append(table_obj)
                
            except Exception as e:
                logger.warning(f"Failed to extract table {table_idx}: {e}")
                continue
        
        return tables
    
    def _extract_images(self, doc: DocxDocument) -> List[Image]:
        """Extract images from DOCX."""
        images = []
        
        try:
            # Get document part
            document_part = doc.part
            
            # Get all image parts
            for rel in document_part.rels.values():
                if "image" in rel.target_ref:
                    try:
                        # Get image binary data
                        image_part = rel.target_part
                        image_data = image_part.blob
                        
                        # Encode to base64
                        base64_data = base64.b64encode(image_data).decode('utf-8')
                        
                        image = Image(
                            image_id=f"docx_img_{len(images) + 1}",
                            base64_data=base64_data,
                            page_number=1  # Can't determine actual page easily
                        )
                        
                        images.append(image)
                        
                    except Exception as e:
                        logger.warning(f"Failed to extract image: {e}")
                        continue
                        
        except Exception as e:
            logger.warning(f"Failed to extract images from DOCX: {e}")
        
        return images
    
    def _analyze_paragraph(self, paragraph) -> tuple[str, Optional[int]]:
        """Analyze paragraph to determine type and hierarchy."""
        text = paragraph.text.strip()
        
        # Check if it's a heading
        hierarchy_level = None
        section_type = "content"
        
        # Check paragraph style for headings
        if paragraph.style.name.startswith('Heading'):
            section_type = "title"
            try:
                # Extract heading level from style name
                hierarchy_level = int(paragraph.style.name.split()[-1])
            except:
                hierarchy_level = 1
        
        # Heuristic checks for titles
        elif (len(text) < 100 and 
              (text.startswith(('第', '一', '二', '三', '四', '五', '六', '七', '八', '九', '十')) or
               any(char.isdigit() for char in text[:10]) and '.' in text[:10])):
            section_type = "title"
            hierarchy_level = 1
        
        # Check for table/image captions
        elif text.lower().startswith(('表', '图', 'table', 'figure')):
            if '表' in text.lower():
                section_type = "table_caption"
            else:
                section_type = "image_caption"
        
        return section_type, hierarchy_level
    
    def _get_raw_text(self, text_content: List[TextContent]) -> str:
        """Get raw text from text content."""
        return "\n".join([content.content for content in text_content])
    
    def _parse_doc_with_antiword(self, file_path: Path) -> StandardizedDocument:
        """
        Parse .doc file using antiword command.
        
        This is a fallback method for old binary .doc format that python-docx cannot read.
        """
        if not ANTIWORD_AVAILABLE:
            raise DocumentProcessingError(
                f"无法解析 .doc 文件 '{file_path.name}'，需要 antiword 工具支持。\n\n"
                f"解决方案：\n"
                f"  方案1（推荐）：将文件转换为 .docx 格式\n"
                f"  方案2（Linux）：安装 antiword 工具\n"
                f"    - Ubuntu/Debian: sudo apt-get install antiword\n"
                f"    - CentOS/RHEL: sudo yum install antiword\n"
                f"    - macOS: brew install antiword\n\n"
                f"注意：antiword 只能提取文本，无法提取图片和表格，建议优先使用方案1。"
            )
        
        try:
            logger.info(f"使用 antiword 解析 .doc 文件: {file_path}")
            
            # 使用 antiword 提取文本
            result = subprocess.run(
                ['antiword', str(file_path)],
                capture_output=True,
                text=True,
                timeout=60
            )
            
            if result.returncode != 0:
                error_msg = result.stderr.strip()
                
                # 检查是否是格式不支持的错误
                if "is not a Word Document" in error_msg:
                    raise DocumentProcessingError(
                        f"无法解析 .doc 文件 '{file_path.name}'：antiword 无法识别该文件格式。\n\n"
                        f"可能的原因：\n"
                        f"1. 该文件不是真正的 Word .doc 格式\n"
                        f"2. 该文件可能损坏\n"
                        f"3. 该文件可能是其他格式（如 RTF、WPS 等）但扩展名为 .doc\n\n"
                        f"建议解决方案：\n"
                        f"1. 使用 Microsoft Word 或 WPS 打开该文件，然后另存为 .docx 格式\n"
                        f"2. 检查文件是否完整、未损坏\n"
                        f"3. 确认文件的真实格式是否为 Word .doc"
                    )
                else:
                    raise DocumentProcessingError(
                        f"antiword 执行失败: {error_msg}"
                    )
            
            raw_text = result.stdout
            
            # 创建 DocumentInfo
            document_info = DocumentInfo(
                file_name=file_path.name,
                file_path=str(file_path.absolute()),
                total_pages=max(1, len(raw_text) // 2000),  # 粗略估算页数
                file_size=self._get_file_size(file_path),
                processing_timestamp=datetime.now()
            )
            
            # 将文本分段处理
            text_content = self._parse_raw_text_to_content(raw_text)
            
            return StandardizedDocument(
                document_info=document_info,
                text_content=text_content,
                tables=[],  # antiword 不提取表格
                images=[],  # antiword 不提取图片
                raw_content=raw_text
            )
            
        except subprocess.TimeoutExpired:
            raise DocumentProcessingError(f"antiword 处理超时: {file_path}")
        except Exception as e:
            raise DocumentProcessingError(f".doc 文件解析失败: {e}")
    
    def _parse_raw_text_to_content(self, raw_text: str) -> List[TextContent]:
        """
        将 antiword 提取的原始文本解析为 TextContent 列表。
        
        Args:
            raw_text: antiword 输出的原始文本
            
        Returns:
            TextContent 对象列表
        """
        text_content = []
        
        # 按段落分割（连续两个换行符）
        paragraphs = raw_text.split('\n\n')
        current_page = 1
        chars_count = 0
        
        for para_idx, paragraph in enumerate(paragraphs):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            
            # 简单的标题检测
            section_type = "content"
            hierarchy_level = None
            
            if len(paragraph) < 100:
                # 可能是标题
                if any(paragraph.startswith(prefix) for prefix in ['第', '一', '二', '三', '四', '五', '1', '2', '3', '4', '5']):
                    section_type = "title"
                    hierarchy_level = 1
            
            content = TextContent(
                section_type=section_type,
                content=paragraph,
                page_number=current_page,
                hierarchy_level=hierarchy_level,
                word_count=len(paragraph.strip())
            )
            
            text_content.append(content)
            
            # 粗略的页数估算（每2000字符约一页）
            chars_count += len(paragraph)
            if chars_count >= 2000:
                current_page += 1
                chars_count = 0
        
        logger.info(f"从 .doc 文件提取了 {len(text_content)} 个段落")
        return text_content